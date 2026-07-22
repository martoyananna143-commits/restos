"""Web data routes: criterion sets, criteria, evaluation types, Google Sheets."""

import logging
import tempfile
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status

from app.api.deps import (
    get_criterion_service,
    get_criterion_set_service,
    get_employee_service,
    get_evaluation_type_service,
    get_google_criterion_sync_service,
    get_google_drive_service,
)
from app.api.routers.web_auth import get_current_web_employee_pin_fresh
from app.internal.services.criterion_service import CriterionService
from app.internal.services.criterion_set_service import CriterionSetService
from app.internal.services.employee_service import EmployeeService
from app.internal.services.evaluation_type_service import EvaluationTypeService
from app.internal.services.google_criterion_sync_service import GoogleCriterionSyncService
from app.internal.services.google_drive_service import GoogleDriveService
from app.internal.services.google_sheet_service import GoogleSheetError
from .helpers import _criterion_set_out, _require_admin
from .schemas import *

logger = logging.getLogger(__name__)

router = APIRouter()

@router.get("/evaluation-types", response_model=list[EvaluationTypeOut])
async def list_evaluation_types(
    emp=Depends(get_current_web_employee_pin_fresh),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    types = await service.get_all(organization_id=emp.organization_id)
    return [
        EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=getattr(t, "description", None))
        for t in types
    ]


@router.post("/evaluation-types", response_model=EvaluationTypeOut, status_code=status.HTTP_201_CREATED)
async def create_evaluation_type(
    data: EvaluationTypeCreateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.evaluation_type.dto import CreateEvaluationTypeDTO
    dto = CreateEvaluationTypeDTO(
        organization_id=emp.organization_id,
        name=data.name,
        code=data.code,
        description=data.description,
    )
    t = await service.create(dto)
    return EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=t.description)


@router.put("/evaluation-types/{type_id}", response_model=EvaluationTypeOut)
async def update_evaluation_type(
    type_id: int,
    data: EvaluationTypeUpdateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await service.get_by_id(type_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Evaluation type not found")

    from app.infra.database.repository.evaluation_type.dto import UpdateEvaluationTypeDTO
    dto = UpdateEvaluationTypeDTO(
        name=data.name,
        code=data.code,
        description=data.description,
    )
    t = await service.update(type_id, dto)
    if not t:
        raise HTTPException(status_code=500, detail="Update failed")
    return EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=t.description)


# ── Criterion sets ────────────────────────────────────────────────────────────

@router.get("/criterion-sets", response_model=list[CriterionSetOut])
async def list_criterion_sets(
    full: bool = False,
    emp=Depends(get_current_web_employee_pin_fresh),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    sets = await service.get_by_organization_id(emp.organization_id)
    return [_criterion_set_out(s) for s in sets]


@router.get("/criterion-sets/{set_id}", response_model=CriterionSetOut)
async def get_criterion_set(
    set_id: int,
    emp=Depends(get_current_web_employee_pin_fresh),
    service: CriterionSetService = Depends(get_criterion_set_service),
    google_sync: GoogleCriterionSyncService = Depends(get_google_criterion_sync_service),
):
    s = await service.get_by_id(set_id)
    if not s or s.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")
    if (getattr(s, "source_type", None) or "internal") != "internal":
        try:
            await google_sync.sync_set(set_id, emp.organization_id, force=False)
            s = await service.get_by_id(set_id) or s
        except GoogleSheetError:
            pass  # return last known state; last_error stored in source_meta
    return _criterion_set_out(s)


@router.post("/criterion-sets", response_model=CriterionSetOut, status_code=status.HTTP_201_CREATED)
async def create_criterion_set(
    data: CriterionSetCreateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO
    dto = CreateCriterionSetDTO(
        organization_id=emp.organization_id,
        name=data.name,
        description=data.description,
        is_default=data.is_default,
        criterion_ids=data.criterion_ids or [],
    )
    s = await service.create(dto)
    return _criterion_set_out(s)


@router.put("/criterion-sets/{set_id}", response_model=CriterionSetOut)
async def update_criterion_set(
    set_id: int,
    data: CriterionSetUpdateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await service.get_by_id(set_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")

    from app.infra.database.repository.criterion_set.dto import UpdateCriterionSetDTO
    dto = UpdateCriterionSetDTO(
        name=data.name,
        description=data.description,
        is_default=data.is_default,
        criterion_ids=data.criterion_ids,
    )
    s = await service.update(set_id, dto)
    if not s:
        raise HTTPException(status_code=500, detail="Update failed")
    return _criterion_set_out(s)


@router.post("/criterion-sets/google/preview", response_model=GoogleSheetPreviewOut)
async def preview_google_criterion_set(
    data: GoogleSheetPreviewIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    google_sync: GoogleCriterionSyncService = Depends(get_google_criterion_sync_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)
    try:
        parsed = await google_sync.preview(data.url)
    except GoogleSheetError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    sample = [
        GoogleSheetPreviewRowOut(
            block=r.block,
            criterion=r.criterion,
            value_type=r.value_type,
        )
        for r in parsed.rows[:20]
    ]
    return GoogleSheetPreviewOut(
        rows_count=len(parsed.rows),
        blocks=parsed.blocks,
        sample_rows=sample,
        spreadsheet_id=parsed.spreadsheet_id,
    )


@router.post("/criterion-sets/google", response_model=CriterionSetOut, status_code=status.HTTP_201_CREATED)
async def create_google_criterion_set(
    data: GoogleSheetCreateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
    google_sync: GoogleCriterionSyncService = Depends(get_google_criterion_sync_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO
    from app.internal.services.google_sheet_service import parse_spreadsheet_url

    try:
        spreadsheet_id, sheet_gid = parse_spreadsheet_url(data.url)
        await google_sync.preview(data.url)
    except GoogleSheetError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    dto = CreateCriterionSetDTO(
        organization_id=emp.organization_id,
        name=data.name.strip(),
        description="Источник: Google Таблица",
        source_type="google_sheet",
        source_url=data.url.strip(),
        source_meta={
            "spreadsheet_id": spreadsheet_id,
            "sheet_gid": sheet_gid,
        },
    )
    created = await service.create(dto)
    try:
        await google_sync.sync_set(created.id, emp.organization_id, force=True)
    except GoogleSheetError as exc:
        await service.delete(created.id)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    refreshed = await service.get_by_id(created.id)
    return _criterion_set_out(refreshed or created)


@router.post("/criterion-sets/google/folder/browse", response_model=GoogleDriveBrowseOut)
async def browse_google_drive_folder(
    data: GoogleDriveBrowseIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    drive_service: GoogleDriveService = Depends(get_google_drive_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)
    try:
        files = drive_service.list_spreadsheets_in_folder(data.folder_url)
    except GoogleSheetError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    return GoogleDriveBrowseOut(
        files=[
            GoogleDriveFileOut(
                file_id=f.file_id,
                name=f.name,
                modified_time=f.modified_time,
                web_view_link=f.web_view_link,
            )
            for f in files
        ]
    )


@router.post("/criterion-sets/google/from-folder", response_model=CriterionSetOut, status_code=status.HTTP_201_CREATED)
async def create_google_criterion_set_from_folder(
    data: GoogleSheetFromFolderIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
    google_sync: GoogleCriterionSyncService = Depends(get_google_criterion_sync_service),
    drive_service: GoogleDriveService = Depends(get_google_drive_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO

    sheet_url = drive_service.spreadsheet_url(data.file_id)
    try:
        await google_sync.preview(sheet_url)
    except GoogleSheetError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    dto = CreateCriterionSetDTO(
        organization_id=emp.organization_id,
        name=data.name.strip(),
        description="Источник: Google Drive",
        source_type="google_sheet",
        source_url=sheet_url,
        source_meta={
            "spreadsheet_id": data.file_id,
            "folder_url": data.folder_url.strip(),
        },
    )
    created = await service.create(dto)
    try:
        await google_sync.sync_set(created.id, emp.organization_id, force=True)
    except GoogleSheetError as exc:
        await service.delete(created.id)
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    refreshed = await service.get_by_id(created.id)
    return _criterion_set_out(refreshed or created)


@router.post("/criterion-sets/{set_id}/google/sync", response_model=CriterionSetOut)
async def sync_google_criterion_set(
    set_id: int,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
    google_sync: GoogleCriterionSyncService = Depends(get_google_criterion_sync_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)
    s = await service.get_by_id(set_id)
    if not s or s.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")
    if (getattr(s, "source_type", None) or "internal") == "internal":
        raise HTTPException(status_code=400, detail="Набор не привязан к Google Таблице")
    try:
        await google_sync.sync_set(set_id, emp.organization_id, force=True)
    except GoogleSheetError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    refreshed = await service.get_by_id(set_id)
    return _criterion_set_out(refreshed or s)


@router.post("/criterion-sets/upload-excel", response_model=dict)
async def upload_excel_criterion_sets(
    file: UploadFile = File(...),
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
):
    """Import criteria and criterion sets from Excel file.

    Format (4 columns): combined_set_name | sub_set_name | question | type
    Types: bool, str, num (or русские: да/нет, текст, число)
    """
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Upload .xlsx or .xls file")

    _MAX_EXCEL_SIZE = 10 * 1024 * 1024  # 10 MB

    _TYPE_MAPPING = {
        "boolean": "bool", "bool": "bool",
        "string": "str", "str": "str", "text": "str",
        "number": "num", "num": "num", "numeric": "num",
        "integer": "num", "int": "num", "float": "num",
        "да/нет": "bool", "да нет": "bool", "логический": "bool", "булевый": "bool",
        "текст": "str", "строка": "str", "текстовый": "str",
        "число": "num", "числовой": "num", "цифра": "num", "оценка": "num",
    }
    _TYPE_TO_DB = {
        "bool": "boolean", "str": "string", "num": "number",
    }
    _HEADER_KW = [
        "название", "набор", "вопрос", "критерий", "тип",
        "name", "question", "type", "set", "категория", "раздел", "субнабор",
    ]
    NUM_COLUMNS = 4

    content = await file.read()
    if len(content) > _MAX_EXCEL_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 10 MB)")
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)
    try:

        try:
            import zipfile
            import xml.etree.ElementTree as ET

            rows_data: list[list[str]] = []
            with zipfile.ZipFile(tmp_path, "r") as zf:
                shared_strings: list[str] = []
                try:
                    with zf.open("xl/sharedStrings.xml") as f:
                        tree = ET.parse(f)
                        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                        for si in tree.getroot().findall(".//x:si", ns):
                            t = si.find(".//x:t", ns)
                            shared_strings.append(t.text if t is not None and t.text else "")
                except KeyError:
                    pass

                with zf.open("xl/worksheets/sheet1.xml") as f:
                    tree = ET.parse(f)
                    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for row in tree.getroot().findall(".//x:row", ns):
                        cells: dict[int, str] = {}
                        for cell in row.findall(".//x:c", ns):
                            ref = cell.get("r")
                            if not ref:
                                continue
                            col_num = 0
                            for ch in ref:
                                if ch.isalpha():
                                    col_num = col_num * 26 + (ord(ch.upper()) - ord("A") + 1)
                                else:
                                    break
                            col_num -= 1
                            v = cell.find("x:v", ns)
                            val = ""
                            if v is not None and v.text:
                                if cell.get("t") == "s":
                                    idx = int(v.text)
                                    val = shared_strings[idx] if idx < len(shared_strings) else ""
                                else:
                                    val = v.text
                            cells[col_num] = val
                        row_vals = [cells.get(i, "") for i in range(NUM_COLUMNS)]
                        rows_data.append(row_vals)

        except Exception as zip_err:
            logger.warning("ZIP read failed: %s", zip_err)
            try:
                from openpyxl import load_workbook
                wb = load_workbook(tmp_path, read_only=True, data_only=True)
                ws = wb.active
                rows_data = []
                for row_tuple in ws.iter_rows(values_only=True):
                    vals = list(row_tuple) + [""] * NUM_COLUMNS
                    rows_data.append([str(v).strip() if v else "" for v in vals[:NUM_COLUMNS]])
                wb.close()
            except Exception as opx_err:
                logger.error("openpyxl failed: %s", opx_err)
                raise HTTPException(status_code=400, detail="Could not read Excel file")

        start_idx = 0
        if rows_data:
            cols_lower = [str(c).lower().strip() for c in rows_data[0]]
            last_col = cols_lower[-1] if cols_lower else ""
            is_header = (
                (any(kw in cols_lower[0] for kw in _HEADER_KW)
                 and any(kw in cols_lower[1] for kw in _HEADER_KW))
                or (last_col and last_col not in _TYPE_MAPPING
                    and any(kw in last_col for kw in _HEADER_KW))
            )
            if is_header:
                start_idx = 1

        combined_sets: OrderedDict[str, OrderedDict[str, list[dict]]] = OrderedDict()
        for row in rows_data[start_idx:]:
            combined_name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
            sub_name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            question = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            raw_type = str(row[3]).strip().lower() if len(row) > 3 and row[3] else "bool"
            if not combined_name or not sub_name or not question:
                continue
            crit_type = _TYPE_MAPPING.get(raw_type, "bool")
            if combined_name not in combined_sets:
                combined_sets[combined_name] = OrderedDict()
            if sub_name not in combined_sets[combined_name]:
                combined_sets[combined_name][sub_name] = []
            combined_sets[combined_name][sub_name].append({"question": question, "type": crit_type})

        if not combined_sets:
            raise HTTPException(status_code=400, detail="No valid data in file")

        from app.infra.database.repository.criterion.dto import CreateCriterionDTO
        from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO

        all_org_criteria = await criterion_service.get_by_organization_id(
            emp.organization_id
        )
        org_by_key: dict[tuple[str, str], int] = {}
        for c in all_org_criteria:
            k = (c.name, c.value_type)
            if k not in org_by_key:
                org_by_key[k] = c.id

        existing_sets_list = await criterion_set_service.get_by_organization_id(
            emp.organization_id, include_inactive=False
        )
        sets_by_name = {s.name: s for s in existing_sets_list}

        touched_set_names: list[str] = []
        created_criteria_count = 0
        global_sort_order = max((c.sort_order for c in all_org_criteria), default=-1) + 1

        for combined_name, sub_sets in combined_sets.items():
            combined_criterion_ids: list[int] = []
            for sub_name, criteria_list in sub_sets.items():
                existing_sub = sets_by_name.get(sub_name)
                in_set_key_to_id: dict[tuple[str, str], int] = {}
                if existing_sub and existing_sub.criterion_ids:
                    crits_in_set = await criterion_service.get_by_ids(
                        existing_sub.criterion_ids
                    )
                    in_set_key_to_id = {
                        (c.name, c.value_type): c.id for c in crits_in_set
                    }

                resolved_this_sub: dict[tuple[str, str], int] = {}
                sub_criterion_ids: list[int] = []
                for idx, crit in enumerate(criteria_list):
                    db_type = _TYPE_TO_DB.get(crit["type"], "boolean")
                    name = str(crit["question"]).strip()
                    key = (name, db_type)
                    if key in resolved_this_sub:
                        cid = resolved_this_sub[key]
                    elif key in in_set_key_to_id:
                        cid = in_set_key_to_id[key]
                        resolved_this_sub[key] = cid
                    elif key in org_by_key:
                        cid = org_by_key[key]
                        resolved_this_sub[key] = cid
                    else:
                        code = f"{sub_name.lower().replace(' ', '_')}_{idx + 1}_{crit['type']}"
                        dto = CreateCriterionDTO(
                            organization_id=emp.organization_id,
                            name=name,
                            code=code,
                            value_type=db_type,
                            description=None,
                            sort_order=global_sort_order,
                        )
                        c = await criterion_service.create(dto)
                        cid = c.id
                        org_by_key[key] = cid
                        resolved_this_sub[key] = cid
                        created_criteria_count += 1
                        global_sort_order += 1
                    sub_criterion_ids.append(cid)
                    combined_criterion_ids.append(cid)

                sub_merged = list(dict.fromkeys(sub_criterion_ids))
                if not sub_merged:
                    continue

                if existing_sub:
                    prev_ids = set(existing_sub.criterion_ids or [])
                    to_add = [i for i in sub_merged if i not in prev_ids]
                    if to_add:
                        await criterion_set_service.append_criteria_to_set(
                            existing_sub.id, to_add
                        )
                    refreshed = await criterion_set_service.get_by_id(existing_sub.id)
                    if refreshed:
                        sets_by_name[sub_name] = refreshed
                    if to_add:
                        touched_set_names.append(sub_name)
                else:
                    sub_dto = CreateCriterionSetDTO(
                        organization_id=emp.organization_id,
                        name=sub_name,
                        description=f"Субнабор из «{combined_name}»",
                        is_default=False,
                        criterion_ids=sub_merged,
                    )
                    created = await criterion_set_service.create(sub_dto)
                    sets_by_name[sub_name] = created
                    touched_set_names.append(sub_name)

            combined_merged = list(dict.fromkeys(combined_criterion_ids))
            if not combined_merged:
                continue

            existing_combined = sets_by_name.get(combined_name)
            if existing_combined:
                prev_c = set(existing_combined.criterion_ids or [])
                to_add_c = [i for i in combined_merged if i not in prev_c]
                if to_add_c:
                    await criterion_set_service.append_criteria_to_set(
                        existing_combined.id, to_add_c
                    )
                refreshed_c = await criterion_set_service.get_by_id(
                    existing_combined.id
                )
                if refreshed_c:
                    sets_by_name[combined_name] = refreshed_c
                if to_add_c:
                    touched_set_names.append(combined_name)
            else:
                combined_dto = CreateCriterionSetDTO(
                    organization_id=emp.organization_id,
                    name=combined_name,
                    description=f"Объединённый набор из: {', '.join(sub_sets.keys())}",
                    is_default=False,
                    criterion_ids=combined_merged,
                )
                created_c = await criterion_set_service.create(combined_dto)
                sets_by_name[combined_name] = created_c
                touched_set_names.append(combined_name)

        touched_unique = list(dict.fromkeys(touched_set_names))

        return {
            "success": True,
            "created_sets": len(touched_unique),
            "created_criteria": created_criteria_count,
            "set_names": touched_unique,
        }
    finally:
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except Exception:
                pass


# ── Criteria (all for org) ────────────────────────────────────────────────────

@router.get("/criteria-all", response_model=list[CriterionOut])
async def list_all_criteria(
    emp=Depends(get_current_web_employee_pin_fresh),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    """List all criteria for the current organization."""
    criteria = await criterion_service.get_by_organization_id(emp.organization_id)
    return [
        CriterionOut(
            id=c.id,
            name=c.name,
            code=c.code,
            description=c.description,
            value_type=getattr(c, "value_type", "boolean"),
            is_required=getattr(c, "is_required", True),
        )
        for c in criteria
    ]


@router.post("/criteria", response_model=CriterionOut, status_code=status.HTTP_201_CREATED)
async def create_criterion(
    data: CriterionCreateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion.dto import CreateCriterionDTO
    dto = CreateCriterionDTO(
        organization_id=emp.organization_id,
        name=data.name,
        code=data.code,
        description=data.description,
        value_type=data.value_type,
        is_required=data.is_required,
    )
    c = await criterion_service.create(dto)
    return CriterionOut(
        id=c.id,
        name=c.name,
        code=c.code,
        description=c.description,
        value_type=c.value_type,
        is_required=c.is_required,
    )


@router.put("/criteria/{criterion_id}", response_model=CriterionOut)
async def update_criterion(
    criterion_id: int,
    data: CriterionUpdateIn,
    emp=Depends(get_current_web_employee_pin_fresh),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await criterion_service.get_by_id(criterion_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion not found")

    from app.infra.database.repository.criterion.dto import UpdateCriterionDTO
    dto = UpdateCriterionDTO(
        name=data.name,
        code=data.code,
        description=data.description,
        value_type=data.value_type,
        is_required=data.is_required,
    )
    c = await criterion_service.update(criterion_id, dto)
    if not c:
        raise HTTPException(status_code=500, detail="Update failed")
    return CriterionOut(
        id=c.id,
        name=c.name,
        code=c.code,
        description=c.description,
        value_type=c.value_type,
        is_required=c.is_required,
    )


@router.get("/criteria", response_model=list[CriterionOut])
async def list_criteria(
    set_id: int,
    emp=Depends(get_current_web_employee_pin_fresh),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
):
    crit_set = await criterion_set_service.get_by_id(set_id)
    if not crit_set or crit_set.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")
    criterion_ids = crit_set.criterion_ids or []
    criteria = await criterion_service.get_by_ids(criterion_ids)
    return [
        CriterionOut(
            id=c.id,
            name=c.name,
            code=c.code,
            description=getattr(c, "description", None),
            value_type=getattr(c, "value_type", "boolean"),
            is_required=getattr(c, "is_required", True),
        )
        for c in criteria
    ]
