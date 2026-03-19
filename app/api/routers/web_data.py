"""Web data router — all data endpoints for the standalone web app.

All endpoints require a valid JWT Bearer token (issued by /api/web/auth/login).
"""

import logging
import tempfile
from collections import OrderedDict
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel

logger = logging.getLogger(__name__)

from app.api.deps import (
    get_criterion_service,
    get_criterion_set_service,
    get_criterion_value_repository,
    get_employee_service,
    get_evaluation_service,
    get_evaluation_type_service,
    get_organization_service,
)
from app.api.routers.web_auth import get_current_web_employee
from app.infra.database.repository.criterion_value.dto import CreateCriterionValueDTO
from app.infra.database.repository.evaluation.dto import (
    CreateEvaluationDTO,
    UpdateEvaluationDTO,
)
from app.internal.usecases.criterion_service import CriterionService
from app.internal.usecases.criterion_set_service import CriterionSetService
from app.internal.usecases.employee_service import EmployeeService
from app.internal.usecases.evaluation_service import EvaluationService
from app.internal.usecases.evaluation_type_service import EvaluationTypeService
from app.internal.usecases.organization_service import OrganizationService
from app.infra.database.repository.criterion_value.criterion_value_asyncpg import (
    CriterionValueRepositoryAsyncpg,
)

router = APIRouter(prefix="/web", tags=["web-data"])


# ── Response schemas ─────────────────────────────────────────────────────────

class EmployeeOut(BaseModel):
    id: int
    full_name: str
    position: Optional[str] = None
    employee_type_id: int
    is_admin: bool = False
    is_active: bool = True


class EmployeeTypeOut(BaseModel):
    id: int
    name: str
    code: str
    is_administrator: bool = False


class EvaluationTypeOut(BaseModel):
    id: int
    name: str
    code: Optional[str] = None
    description: Optional[str] = None


class CriterionSetOut(BaseModel):
    id: int
    name: str
    description: Optional[str] = None
    is_default: bool = False
    criterion_ids: Optional[list[int]] = None


class CriterionSetCreateIn(BaseModel):
    name: str
    description: Optional[str] = None
    is_default: bool = False
    criterion_ids: list[int] = []


class CriterionSetUpdateIn(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    is_default: Optional[bool] = None
    criterion_ids: Optional[list[int]] = None


class CriterionCreateIn(BaseModel):
    name: str
    code: str
    description: Optional[str] = None
    value_type: str = "boolean"  # boolean, string, number
    is_required: bool = True


class CriterionUpdateIn(BaseModel):
    name: Optional[str] = None
    code: Optional[str] = None
    description: Optional[str] = None
    value_type: Optional[str] = None
    is_required: Optional[bool] = None


class CriterionOut(BaseModel):
    id: int
    name: str
    code: str
    description: Optional[str] = None
    value_type: str = "boolean"
    is_required: bool = True


class AnalyticsOut(BaseModel):
    total_evaluations: int
    average_score: float
    employees_count: int
    top_employees: list[dict]
    criteria_stats: list[dict]


class EvaluationOut(BaseModel):
    id: int
    evaluated_employee_id: Optional[int] = None
    evaluated_employee_name: Optional[str] = None
    evaluation_type_id: int
    evaluation_type_name: Optional[str] = None
    criterion_set_id: int
    score_percentage: Optional[float] = None
    status: str = "completed"
    created_at: str


class FormDataOut(BaseModel):
    criteria: list[CriterionOut]
    organization_name: str
    evaluated_employee_name: Optional[str] = None
    criterion_set_name: Optional[str] = None
    evaluation_id: int


class AnswerIn(BaseModel):
    criterion_id: int
    value: bool | str | float
    comment: Optional[str] = None


class StartEvaluationRequest(BaseModel):
    evaluated_employee_id: int
    criterion_set_id: int
    evaluation_type_id: int


class StartEvaluationResponse(BaseModel):
    evaluation_id: int
    criteria: list[CriterionOut]
    organization_name: str
    evaluated_employee_name: str
    criterion_set_name: str


class SubmitEvaluationRequest(BaseModel):
    answers: list[AnswerIn]
    comment: Optional[str] = None


class SubmitEvaluationResponse(BaseModel):
    evaluation_id: int
    score_percentage: float
    passed_criteria: int
    failed_criteria: int
    total_criteria: int


class UpdateRoleRequest(BaseModel):
    employee_type_id: int


# ── Helpers ──────────────────────────────────────────────────────────────────

def _require_same_org(current_emp, target_org_id: int):
    if current_emp.organization_id != target_org_id:
        raise HTTPException(status_code=403, detail="Access denied")


def _require_admin(emp, employee_types: list[dict]):
    my_type = next((t for t in employee_types if t["id"] == emp.employee_type_id), {})
    if not my_type.get("is_administrator", False):
        raise HTTPException(status_code=403, detail="Admin access required")


# ── Employees ────────────────────────────────────────────────────────────────

@router.get("/employees", response_model=list[EmployeeOut])
async def list_employees(
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
):
    employees = await employee_service.get_by_organization_id(emp.organization_id)
    emp_types = await employee_service.get_employee_types()
    type_map = {t["id"]: t for t in emp_types}
    return [
        EmployeeOut(
            id=e.id,
            full_name=e.full_name,
            position=e.position,
            employee_type_id=e.employee_type_id,
            is_admin=bool(type_map.get(e.employee_type_id, {}).get("is_administrator", False)),
            is_active=e.is_active,
        )
        for e in employees
    ]


@router.get("/employee-types", response_model=list[EmployeeTypeOut])
async def list_employee_types(
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
):
    types = await employee_service.get_employee_types()
    return [
        EmployeeTypeOut(
            id=t["id"],
            name=t["name"],
            code=t["code"],
            is_administrator=bool(t.get("is_administrator", False)),
        )
        for t in types
    ]


@router.put("/employees/{employee_id}/role")
async def update_employee_role(
    employee_id: int,
    data: UpdateRoleRequest,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    target = await employee_service.get_by_id(employee_id)
    if not target or target.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Employee not found")

    from app.infra.database.repository.employee.dto import UpdateEmployeeDTO
    updated = await employee_service.update(employee_id, UpdateEmployeeDTO(employee_type_id=data.employee_type_id))
    if not updated:
        raise HTTPException(status_code=500, detail="Update failed")
    return {"success": True}


# ── Evaluation types ─────────────────────────────────────────────────────────

class EvaluationTypeCreateIn(BaseModel):
    name: str
    code: str
    description: Optional[str] = None


class EvaluationTypeUpdateIn(BaseModel):
    name: Optional[str] = None
    code: Optional[str] = None
    description: Optional[str] = None


@router.get("/evaluation-types", response_model=list[EvaluationTypeOut])
async def list_evaluation_types(
    emp=Depends(get_current_web_employee),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    types = await service.get_all(organization_id=emp.organization_id)
    return [
        EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=getattr(t, "description", None))
        for t in types
    ]


@router.post("/evaluation-types", response_model=EvaluationTypeOut, status_code=status.HTTP_201_CREATED)
async def create_evaluation_type(
    data: EvaluationTypeCreateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.evaluation_type.dto import CreateEvaluationTypeDTO
    dto = CreateEvaluationTypeDTO(
        organization_id=emp.organization_id,
        name=data.name,
        code=data.code,
        description=data.description,
    )
    t = await service.create(dto)
    return EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=t.description)


@router.put("/evaluation-types/{type_id}", response_model=EvaluationTypeOut)
async def update_evaluation_type(
    type_id: int,
    data: EvaluationTypeUpdateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await service.get_by_id(type_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Evaluation type not found")

    from app.infra.database.repository.evaluation_type.dto import UpdateEvaluationTypeDTO
    dto = UpdateEvaluationTypeDTO(
        name=data.name,
        code=data.code,
        description=data.description,
    )
    t = await service.update(type_id, dto)
    if not t:
        raise HTTPException(status_code=500, detail="Update failed")
    return EvaluationTypeOut(id=t.id, name=t.name, code=t.code, description=t.description)


# ── Criterion sets ────────────────────────────────────────────────────────────

@router.get("/criterion-sets", response_model=list[CriterionSetOut])
async def list_criterion_sets(
    full: bool = False,
    emp=Depends(get_current_web_employee),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    sets = await service.get_by_organization_id(emp.organization_id)
    return [
        CriterionSetOut(
            id=s.id,
            name=s.name,
            description=getattr(s, "description", None),
            is_default=s.is_default,
            # Always include criterion_ids so the select phase can show counts;
            # the `full` flag is kept for backwards-compat but no longer hides IDs.
            criterion_ids=s.criterion_ids or [],
        )
        for s in sets
    ]


@router.get("/criterion-sets/{set_id}", response_model=CriterionSetOut)
async def get_criterion_set(
    set_id: int,
    emp=Depends(get_current_web_employee),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    s = await service.get_by_id(set_id)
    if not s or s.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")
    return CriterionSetOut(
        id=s.id,
        name=s.name,
        description=s.description,
        is_default=s.is_default,
        criterion_ids=s.criterion_ids or [],
    )


@router.post("/criterion-sets", response_model=CriterionSetOut, status_code=status.HTTP_201_CREATED)
async def create_criterion_set(
    data: CriterionSetCreateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO
    dto = CreateCriterionSetDTO(
        organization_id=emp.organization_id,
        name=data.name,
        description=data.description,
        is_default=data.is_default,
        criterion_ids=data.criterion_ids or [],
    )
    s = await service.create(dto)
    return CriterionSetOut(
        id=s.id,
        name=s.name,
        description=s.description,
        is_default=s.is_default,
        criterion_ids=s.criterion_ids or [],
    )


@router.put("/criterion-sets/{set_id}", response_model=CriterionSetOut)
async def update_criterion_set(
    set_id: int,
    data: CriterionSetUpdateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    service: CriterionSetService = Depends(get_criterion_set_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await service.get_by_id(set_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion set not found")

    from app.infra.database.repository.criterion_set.dto import UpdateCriterionSetDTO
    dto = UpdateCriterionSetDTO(
        name=data.name,
        description=data.description,
        is_default=data.is_default,
        criterion_ids=data.criterion_ids,
    )
    s = await service.update(set_id, dto)
    if not s:
        raise HTTPException(status_code=500, detail="Update failed")
    return CriterionSetOut(
        id=s.id,
        name=s.name,
        description=s.description,
        is_default=s.is_default,
        criterion_ids=s.criterion_ids or [],
    )


@router.post("/criterion-sets/upload-excel", response_model=dict)
async def upload_excel_criterion_sets(
    file: UploadFile = File(...),
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
):
    """Import criteria and criterion sets from Excel file.

    Format (4 columns): combined_set_name | sub_set_name | question | type
    Types: bool, str, num (or русские: да/нет, текст, число)
    """
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Upload .xlsx or .xls file")

    _TYPE_MAPPING = {
        "boolean": "bool", "bool": "bool",
        "string": "str", "str": "str", "text": "str",
        "number": "num", "num": "num", "numeric": "num",
        "integer": "num", "int": "num", "float": "num",
        "да/нет": "bool", "да нет": "bool", "логический": "bool", "булевый": "bool",
        "текст": "str", "строка": "str", "текстовый": "str",
        "число": "num", "числовой": "num", "цифра": "num", "оценка": "num",
    }
    _TYPE_TO_DB = {
        "bool": "boolean", "str": "string", "num": "number",
    }
    _HEADER_KW = [
        "название", "набор", "вопрос", "критерий", "тип",
        "name", "question", "type", "set", "категория", "раздел", "субнабор",
    ]
    NUM_COLUMNS = 4

    content = await file.read()
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp.write(content)
        tmp_path = Path(tmp.name)
    try:

        try:
            import zipfile
            import xml.etree.ElementTree as ET

            rows_data: list[list[str]] = []
            with zipfile.ZipFile(tmp_path, "r") as zf:
                shared_strings: list[str] = []
                try:
                    with zf.open("xl/sharedStrings.xml") as f:
                        tree = ET.parse(f)
                        ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                        for si in tree.getroot().findall(".//x:si", ns):
                            t = si.find(".//x:t", ns)
                            shared_strings.append(t.text if t is not None and t.text else "")
                except KeyError:
                    pass

                with zf.open("xl/worksheets/sheet1.xml") as f:
                    tree = ET.parse(f)
                    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for row in tree.getroot().findall(".//x:row", ns):
                        cells: dict[int, str] = {}
                        for cell in row.findall(".//x:c", ns):
                            ref = cell.get("r")
                            if not ref:
                                continue
                            col_num = 0
                            for ch in ref:
                                if ch.isalpha():
                                    col_num = col_num * 26 + (ord(ch.upper()) - ord("A") + 1)
                                else:
                                    break
                            col_num -= 1
                            v = cell.find("x:v", ns)
                            val = ""
                            if v is not None and v.text:
                                if cell.get("t") == "s":
                                    idx = int(v.text)
                                    val = shared_strings[idx] if idx < len(shared_strings) else ""
                                else:
                                    val = v.text
                            cells[col_num] = val
                        row_vals = [cells.get(i, "") for i in range(NUM_COLUMNS)]
                        rows_data.append(row_vals)

        except Exception as zip_err:
            logger.warning("ZIP read failed: %s", zip_err)
            try:
                from openpyxl import load_workbook
                wb = load_workbook(tmp_path, read_only=True, data_only=True)
                ws = wb.active
                rows_data = []
                for row_tuple in ws.iter_rows(values_only=True):
                    vals = list(row_tuple) + [""] * NUM_COLUMNS
                    rows_data.append([str(v).strip() if v else "" for v in vals[:NUM_COLUMNS]])
                wb.close()
            except Exception as opx_err:
                logger.error("openpyxl failed: %s", opx_err)
                raise HTTPException(status_code=400, detail="Could not read Excel file")

        start_idx = 0
        if rows_data:
            cols_lower = [str(c).lower().strip() for c in rows_data[0]]
            last_col = cols_lower[-1] if cols_lower else ""
            is_header = (
                (any(kw in cols_lower[0] for kw in _HEADER_KW)
                 and any(kw in cols_lower[1] for kw in _HEADER_KW))
                or (last_col and last_col not in _TYPE_MAPPING
                    and any(kw in last_col for kw in _HEADER_KW))
            )
            if is_header:
                start_idx = 1

        combined_sets: OrderedDict[str, OrderedDict[str, list[dict]]] = OrderedDict()
        for row in rows_data[start_idx:]:
            combined_name = str(row[0]).strip() if len(row) > 0 and row[0] else ""
            sub_name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            question = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            raw_type = str(row[3]).strip().lower() if len(row) > 3 and row[3] else "bool"
            if not combined_name or not sub_name or not question:
                continue
            crit_type = _TYPE_MAPPING.get(raw_type, "bool")
            if combined_name not in combined_sets:
                combined_sets[combined_name] = OrderedDict()
            if sub_name not in combined_sets[combined_name]:
                combined_sets[combined_name][sub_name] = []
            combined_sets[combined_name][sub_name].append({"question": question, "type": crit_type})

        if not combined_sets:
            raise HTTPException(status_code=400, detail="No valid data in file")

        from app.infra.database.repository.criterion.dto import CreateCriterionDTO
        from app.infra.database.repository.criterion_set.dto import CreateCriterionSetDTO

        created_sets: list[str] = []
        created_criteria_count = 0
        global_sort_order = 0

        for combined_name, sub_sets in combined_sets.items():
            combined_criterion_ids: list[int] = []
            for sub_name, criteria_list in sub_sets.items():
                sub_criterion_ids: list[int] = []
                for idx, crit in enumerate(criteria_list):
                    db_type = _TYPE_TO_DB.get(crit["type"], "boolean")
                    code = f"{sub_name.lower().replace(' ', '_')}_{idx + 1}_{crit['type']}"
                    dto = CreateCriterionDTO(
                        organization_id=emp.organization_id,
                        name=crit["question"],
                        code=code,
                        value_type=db_type,
                        description=None,
                        sort_order=global_sort_order,
                    )
                    c = await criterion_service.create(dto)
                    sub_criterion_ids.append(c.id)
                    combined_criterion_ids.append(c.id)
                    created_criteria_count += 1
                    global_sort_order += 1

                if sub_criterion_ids:
                    sub_dto = CreateCriterionSetDTO(
                        organization_id=emp.organization_id,
                        name=sub_name,
                        description=f"Субнабор из «{combined_name}»",
                        is_default=False,
                        criterion_ids=sub_criterion_ids,
                    )
                    await criterion_set_service.create(sub_dto)
                    created_sets.append(sub_name)

            if combined_criterion_ids:
                combined_dto = CreateCriterionSetDTO(
                    organization_id=emp.organization_id,
                    name=combined_name,
                    description=f"Объединённый набор из: {', '.join(sub_sets.keys())}",
                    is_default=False,
                    criterion_ids=combined_criterion_ids,
                )
                await criterion_set_service.create(combined_dto)
                created_sets.append(combined_name)

        return {
            "success": True,
            "created_sets": len(created_sets),
            "created_criteria": created_criteria_count,
            "set_names": created_sets,
        }
    finally:
        if tmp_path.exists():
            try:
                tmp_path.unlink()
            except Exception:
                pass


# ── Criteria (all for org) ────────────────────────────────────────────────────

@router.get("/criteria-all", response_model=list[CriterionOut])
async def list_all_criteria(
    emp=Depends(get_current_web_employee),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    """List all criteria for the current organization."""
    criteria = await criterion_service.get_by_organization_id(emp.organization_id)
    return [
        CriterionOut(
            id=c.id,
            name=c.name,
            code=c.code,
            description=c.description,
            value_type=getattr(c, "value_type", "boolean"),
            is_required=getattr(c, "is_required", True),
        )
        for c in criteria
    ]


@router.post("/criteria", response_model=CriterionOut, status_code=status.HTTP_201_CREATED)
async def create_criterion(
    data: CriterionCreateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    from app.infra.database.repository.criterion.dto import CreateCriterionDTO
    dto = CreateCriterionDTO(
        organization_id=emp.organization_id,
        name=data.name,
        code=data.code,
        description=data.description,
        value_type=data.value_type,
        is_required=data.is_required,
    )
    c = await criterion_service.create(dto)
    return CriterionOut(
        id=c.id,
        name=c.name,
        code=c.code,
        description=c.description,
        value_type=c.value_type,
        is_required=c.is_required,
    )


@router.put("/criteria/{criterion_id}", response_model=CriterionOut)
async def update_criterion(
    criterion_id: int,
    data: CriterionUpdateIn,
    emp=Depends(get_current_web_employee),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    emp_types = await employee_service.get_employee_types()
    _require_admin(emp, emp_types)

    existing = await criterion_service.get_by_id(criterion_id)
    if not existing or existing.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Criterion not found")

    from app.infra.database.repository.criterion.dto import UpdateCriterionDTO
    dto = UpdateCriterionDTO(
        name=data.name,
        code=data.code,
        description=data.description,
        value_type=data.value_type,
        is_required=data.is_required,
    )
    c = await criterion_service.update(criterion_id, dto)
    if not c:
        raise HTTPException(status_code=500, detail="Update failed")
    return CriterionOut(
        id=c.id,
        name=c.name,
        code=c.code,
        description=c.description,
        value_type=c.value_type,
        is_required=c.is_required,
    )


@router.get("/criteria", response_model=list[CriterionOut])
async def list_criteria(
    set_id: int,
    emp=Depends(get_current_web_employee),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
):
    crit_set = await criterion_set_service.get_by_id(set_id)
    if not crit_set:
        raise HTTPException(status_code=404, detail="Criterion set not found")
    criterion_ids = crit_set.criterion_ids or []
    criteria = await criterion_service.get_by_ids(criterion_ids)
    return [
        CriterionOut(
            id=c.id,
            name=c.name,
            code=c.code,
            description=getattr(c, "description", None),
            value_type=getattr(c, "value_type", "boolean"),
            is_required=getattr(c, "is_required", True),
        )
        for c in criteria
    ]


# ── Evaluations ───────────────────────────────────────────────────────────────

@router.get("/evaluations", response_model=list[EvaluationOut])
async def list_evaluations(
    emp=Depends(get_current_web_employee),
    evaluation_service: EvaluationService = Depends(get_evaluation_service),
    employee_service: EmployeeService = Depends(get_employee_service),
    eval_type_service: EvaluationTypeService = Depends(get_evaluation_type_service),
):
    evaluations = await evaluation_service.get_by_organization_id(emp.organization_id)

    # Build lookup caches
    eval_types = await eval_type_service.get_all(organization_id=emp.organization_id)
    type_map = {t.id: t.name for t in eval_types}

    all_employees = await employee_service.get_by_organization_id(emp.organization_id)
    emp_map = {e.id: e.full_name for e in all_employees}

    result = []
    for ev in evaluations:
        score = None
        if hasattr(ev, "score_percentage"):
            score = ev.score_percentage
        result.append(EvaluationOut(
            id=ev.id,
            evaluated_employee_id=getattr(ev, "evaluated_employee_id", None),
            evaluated_employee_name=emp_map.get(getattr(ev, "evaluated_employee_id", 0)),
            evaluation_type_id=ev.evaluation_type_id,
            evaluation_type_name=type_map.get(ev.evaluation_type_id),
            criterion_set_id=getattr(ev, "criterion_set_id", 0),
            score_percentage=score,
            status="completed" if score is not None else "in_progress",
            created_at=str(ev.created_at),
        ))
    return result


@router.post("/evaluations/start", response_model=StartEvaluationResponse)
async def start_evaluation(
    data: StartEvaluationRequest,
    emp=Depends(get_current_web_employee),
    evaluation_service: EvaluationService = Depends(get_evaluation_service),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
    org_service: OrganizationService = Depends(get_organization_service),
):
    """Start a new evaluation and return criteria to fill."""
    target_emp = await employee_service.get_by_id(data.evaluated_employee_id)
    if not target_emp or target_emp.organization_id != emp.organization_id:
        raise HTTPException(status_code=404, detail="Target employee not found")

    crit_set = await criterion_set_service.get_by_id(data.criterion_set_id)
    if not crit_set:
        raise HTTPException(status_code=404, detail="Criterion set not found")

    org = await org_service.get_by_id(emp.organization_id)
    criterion_ids = crit_set.criterion_ids or []
    criteria = await criterion_service.get_by_ids(criterion_ids)

    dto = CreateEvaluationDTO(
        organization_id=emp.organization_id,
        criterion_set_id=data.criterion_set_id,
        evaluation_type_id=data.evaluation_type_id,
        filled_by_employee_id=emp.id,
        evaluated_employee_id=data.evaluated_employee_id,
    )
    evaluation = await evaluation_service.create(dto)

    return StartEvaluationResponse(
        evaluation_id=evaluation.id,
        criteria=[
            CriterionOut(
                id=c.id,
                name=c.name,
                code=c.code,
                description=getattr(c, "description", None),
                value_type=getattr(c, "value_type", "boolean"),
                is_required=getattr(c, "is_required", True),
            )
            for c in criteria
        ],
        organization_name=org.name if org else "",
        evaluated_employee_name=target_emp.full_name,
        criterion_set_name=crit_set.name,
    )


@router.post("/evaluations/{evaluation_id}/submit", response_model=SubmitEvaluationResponse)
async def submit_evaluation(
    evaluation_id: int,
    data: SubmitEvaluationRequest,
    emp=Depends(get_current_web_employee),
    evaluation_service: EvaluationService = Depends(get_evaluation_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
    criterion_set_service: CriterionSetService = Depends(get_criterion_set_service),
    cv_repo: CriterionValueRepositoryAsyncpg = Depends(get_criterion_value_repository),
):
    evaluation = await evaluation_service.get_by_id(evaluation_id)
    if not evaluation:
        raise HTTPException(status_code=404, detail="Evaluation not found")
    if evaluation.organization_id != emp.organization_id:
        raise HTTPException(status_code=403, detail="Access denied")

    crit_set = await criterion_set_service.get_by_id(evaluation.criterion_set_id or 0)
    criterion_ids = crit_set.criterion_ids if crit_set else []
    criteria = await criterion_service.get_by_ids(criterion_ids or [])
    crit_map = {c.id: c for c in criteria}

    passed = 0
    failed = 0
    score_points: list[float] = []  # 0.0 – 1.0 per answered criterion

    for answer in data.answers:
        crit = crit_map.get(answer.criterion_id)
        if not crit:
            continue
        value_type = getattr(crit, "value_type", "boolean")

        if value_type == "boolean":
            bool_val = (
                answer.value if isinstance(answer.value, bool)
                else str(answer.value).lower() in ("true", "1")
            )
            if bool_val:
                passed += 1
                score_points.append(1.0)
            else:
                failed += 1
                score_points.append(0.0)
            dto = CreateCriterionValueDTO(
                evaluation_id=evaluation_id,
                criterion_id=answer.criterion_id,
                value=bool_val,
                notes=answer.comment,
            )
        elif value_type == "number":
            try:
                num = float(answer.value) if not isinstance(answer.value, bool) else 0.0
            except (ValueError, TypeError):
                num = 0.0
            # Normalize 1–5 scale to 0–1; clamp to [0, 1]
            norm = max(0.0, min(1.0, (num - 1.0) / 4.0)) if num > 0 else 0.0
            passed += 1
            score_points.append(norm)
            dto = CreateCriterionValueDTO(
                evaluation_id=evaluation_id,
                criterion_id=answer.criterion_id,
                value=num,
                notes=answer.comment,
            )
        else:
            # Text criterion: any non-empty answer counts as full score
            text_val = str(answer.value).strip()
            score_points.append(1.0 if text_val else 0.0)
            passed += 1 if text_val else 0
            dto = CreateCriterionValueDTO(
                evaluation_id=evaluation_id,
                criterion_id=answer.criterion_id,
                value=text_val,
                notes=answer.comment,
            )
        await cv_repo.create(dto)

    total = len(data.answers)
    score_pct = (sum(score_points) / len(score_points) * 100.0) if score_points else 0.0

    update_dto = UpdateEvaluationDTO(
        score_percentage=score_pct,
        total_criteria=total,
        passed_criteria=passed,
        failed_criteria=failed,
        comment=data.comment,
        status="completed",
    )
    await evaluation_service.update(evaluation_id, update_dto)

    return SubmitEvaluationResponse(
        evaluation_id=evaluation_id,
        score_percentage=score_pct,
        passed_criteria=passed,
        failed_criteria=failed,
        total_criteria=total,
    )


# ── Analytics ────────────────────────────────────────────────────────────────

@router.get("/analytics", response_model=AnalyticsOut)
async def get_analytics(
    emp=Depends(get_current_web_employee),
    evaluation_service: EvaluationService = Depends(get_evaluation_service),
    employee_service: EmployeeService = Depends(get_employee_service),
    criterion_service: CriterionService = Depends(get_criterion_service),
):
    org_id = emp.organization_id
    evaluations = await evaluation_service.get_by_organization_id(org_id)
    employees = await employee_service.get_by_organization_id(org_id)

    total = len(evaluations)
    scores = [
        ev.score_percentage
        for ev in evaluations
        if getattr(ev, "score_percentage", None) is not None
    ]
    avg_score = sum(scores) / len(scores) if scores else 0.0

    # Per-employee average
    emp_scores: dict[int, list[float]] = {}
    for ev in evaluations:
        eid = getattr(ev, "evaluated_employee_id", None)
        sc = getattr(ev, "score_percentage", None)
        if eid and sc is not None:
            emp_scores.setdefault(eid, []).append(sc)

    emp_map = {e.id: e.full_name for e in employees}
    top_employees = sorted(
        [
            {"id": eid, "name": emp_map.get(eid, ""), "avg": sum(s) / len(s), "count": len(s)}
            for eid, s in emp_scores.items()
        ],
        key=lambda x: x["avg"],
        reverse=True,
    )[:10]

    return AnalyticsOut(
        total_evaluations=total,
        average_score=round(avg_score, 1),
        employees_count=len(employees),
        top_employees=top_employees,
        criteria_stats=[],
    )
