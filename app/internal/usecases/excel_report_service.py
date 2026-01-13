"""Excel Report Service for generating evaluation reports."""

import logging
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

logger = logging.getLogger(__name__)


class ExcelReportService:
    """Service for generating Excel reports from evaluation data."""

    def __init__(self):
        """Initialize Excel report service."""
        pass

    def generate_evaluation_report(
        self,
        evaluation_id: int,
        evaluation_type_name: str,
        organization_name: str,
        evaluated_employee_name: str,
        filled_by_employee_name: str,
        evaluation_date: str,
        criteria: List[Dict],
        total_criteria: int,
        passed_criteria: int,
        failed_criteria: int,
        score_percentage: float,
        criterion_set_name: Optional[str] = None,
        boolean_criteria_count: Optional[int] = None,
    ) -> Path:
        """Generate Excel report for evaluation.

        Args:
            evaluation_id: Evaluation ID.
            evaluation_type_name: Name of evaluation type.
            organization_name: Name of organization.
            evaluated_employee_name: Name of evaluated employee.
            filled_by_employee_name: Name of employee who filled the evaluation.
            evaluation_date: Date of evaluation.
            criteria: List of criteria with answers and comments.
            total_criteria: Total number of criteria.
            passed_criteria: Number of passed criteria.
            failed_criteria: Number of failed criteria.
            score_percentage: Percentage of successful criteria.
            criterion_set_name: Optional name of criterion set.
            boolean_criteria_count: Number of boolean criteria.

        Returns:
            Path to generated Excel file.
        """
        try:
            temp_dir = Path(tempfile.gettempdir())
            excel_filename = (
                f"evaluation_report_{evaluation_id}_"
                f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            )
            excel_path = temp_dir / excel_filename

            # Создаём рабочую книгу
            wb = Workbook()
            ws = wb.active
            ws.title = "Отчет об оценке"

            # Стили
            title_font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
            header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            normal_font = Font(name="Arial", size=11)
            bold_font = Font(name="Arial", size=11, bold=True)

            # Цвета
            header_fill = PatternFill(
                start_color="4CAF50", end_color="4CAF50", fill_type="solid"
            )
            title_fill = PatternFill(
                start_color="2c3e50", end_color="2c3e50", fill_type="solid"
            )
            green_fill = PatternFill(
                start_color="d4edda", end_color="d4edda", fill_type="solid"
            )
            red_fill = PatternFill(
                start_color="f8d7da", end_color="f8d7da", fill_type="solid"
            )
            light_gray_fill = PatternFill(
                start_color="f9f9f9", end_color="f9f9f9", fill_type="solid"
            )
            summary_fill = PatternFill(
                start_color="e8f5e9", end_color="e8f5e9", fill_type="solid"
            )

            # Границы
            thin_border = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin"),
            )
            thick_border = Border(
                left=Side(style="thick"),
                right=Side(style="thick"),
                top=Side(style="thick"),
                bottom=Side(style="thick"),
            )

            row = 1

            # Заголовок
            ws.merge_cells(f"A{row}:D{row}")
            cell = ws[f"A{row}"]
            cell.value = "Отчет об оценке"
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thick_border
            ws.row_dimensions[row].height = 30
            row += 1

            # Подзаголовок
            ws.merge_cells(f"A{row}:D{row}")
            cell = ws[f"A{row}"]
            cell.value = evaluation_type_name
            cell.font = Font(name="Arial", size=14, color="7f8c8d")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

            ws.merge_cells(f"A{row}:D{row}")
            cell = ws[f"A{row}"]
            cell.value = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            cell.font = Font(name="Arial", size=10, color="7f8c8d")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 2

            # Информационная секция
            info_data = [
                ["Организация:", organization_name],
                ["Оцениваемый сотрудник:", evaluated_employee_name],
                ["Оценивающий сотрудник:", filled_by_employee_name],
            ]

            if criterion_set_name:
                info_data.append(["Набор критериев:", criterion_set_name])

            info_data.append(["Дата оценки:", evaluation_date])

            for label, value in info_data:
                ws[f"A{row}"] = label
                ws[f"A{row}"].font = bold_font
                ws[f"A{row}"].fill = light_gray_fill
                ws[f"A{row}"].border = thin_border
                ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")

                ws[f"B{row}"] = value
                ws[f"B{row}"].font = normal_font
                ws[f"B{row}"].fill = light_gray_fill
                ws[f"B{row}"].border = thin_border
                ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="center")

                ws.merge_cells(f"B{row}:D{row}")
                row += 1

            row += 1

            # Заголовок таблицы критериев
            headers = ["№", "Критерий", "Результат", "Комментарий"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=row, column=col_idx)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thick_border
            ws.row_dimensions[row].height = 25
            row += 1

            # Данные критериев
            for idx, criterion in enumerate(criteria, 1):
                value = criterion.get("value")
                value_type: str = criterion.get("value_type", "boolean")

                # Парсим JSON строки, если значение - строка с JSON
                if isinstance(value, str):
                    value_stripped = value.strip()
                    if value_stripped.startswith('{') and value_stripped.endswith('}'):
                        try:
                            import json
                            parsed = json.loads(value)
                            if isinstance(parsed, dict) and "value" in parsed:
                                # Извлекаем значение из JSON объекта {"type": "...", "value": ...}
                                extracted_value = parsed.get("value")
                                if extracted_value is not None:
                                    # Преобразуем тип в зависимости от type в JSON
                                    value_type_from_json = parsed.get("type")
                                    if value_type_from_json == "boolean":
                                        value = bool(extracted_value) if not isinstance(extracted_value, bool) else extracted_value
                                    elif value_type_from_json == "number":
                                        if isinstance(extracted_value, str):
                                            try:
                                                value = float(extracted_value) if "." in extracted_value else int(extracted_value)
                                            except (ValueError, TypeError):
                                                value = extracted_value
                                        else:
                                            value = extracted_value
                                    elif value_type_from_json == "string":
                                        value = str(extracted_value)
                                    else:
                                        value = extracted_value
                        except (json.JSONDecodeError, TypeError, ValueError):
                            # Если не удалось распарсить, используем как есть
                            pass

                # Номер
                cell = ws.cell(row=row, column=1)
                cell.value = idx
                cell.font = normal_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

                # Название критерия
                cell = ws.cell(row=row, column=2)
                cell.value = criterion.get("name", "Не указано")
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = thin_border

                # Результат
                result_cell = ws.cell(row=row, column=3)
                if value_type == "boolean":
                    if isinstance(value, bool):
                        is_passed = value
                    elif isinstance(value, str):
                        is_passed = value.lower() in ("true", "1", "yes", "да")
                    elif isinstance(value, (int, float)):
                        is_passed = bool(value)
                    else:
                        is_passed = False

                    result_cell.value = "Да" if is_passed else "Нет"
                    result_cell.fill = green_fill if is_passed else red_fill
                    result_cell.font = Font(
                        name="Arial", size=11, bold=True, color="155724" if is_passed else "721c24"
                    )
                elif value_type == "string":
                    result_cell.value = str(value) if value is not None else "—"
                    result_cell.font = normal_font
                elif value_type == "number":
                    if value is None:
                        result_cell.value = "—"
                    else:
                        if isinstance(value, float):
                            result_cell.value = float(value)
                            result_cell.number_format = "0.00"
                        else:
                            result_cell.value = int(value) if isinstance(value, float) and value.is_integer() else value
                    result_cell.font = normal_font
                else:
                    result_cell.value = str(value) if value is not None else "—"
                    result_cell.font = normal_font

                result_cell.alignment = Alignment(horizontal="center", vertical="center")
                result_cell.border = thin_border

                # Комментарий
                cell = ws.cell(row=row, column=4)
                comment = criterion.get("comment", "—")
                if not comment:
                    comment = "—"
                cell.value = comment
                cell.font = Font(name="Arial", size=11, italic=True)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                cell.border = thin_border

                # Чередующийся цвет фона
                if idx % 2 == 0:
                    for col in range(1, 5):
                        ws.cell(row=row, column=col).fill = light_gray_fill

                row += 1

            row += 1

            # Итоговая статистика
            stats_title = ws.cell(row=row, column=1)
            stats_title.value = "Итоговая статистика"
            stats_title.font = Font(name="Arial", size=14, bold=True, color="2c3e50")
            row += 1

            if boolean_criteria_count and boolean_criteria_count > 0:
                # Статистика по булевым критериям
                stats_headers = [
                    "Всего критериев",
                    "Пройдено (булевые)",
                    "Не пройдено (булевые)",
                    "Успешность (булевые)",
                ]
                stats_values = [
                    total_criteria,
                    passed_criteria,
                    failed_criteria,
                    f"{score_percentage:.1f}%",
                ]

                for col_idx, (header, value) in enumerate(zip(stats_headers, stats_values), 1):
                    cell = ws.cell(row=row, column=col_idx)
                    cell.value = str(header)
                    cell.font = Font(name="Arial", size=9)
                    cell.fill = summary_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thin_border

                    value_cell = ws.cell(row=row + 1, column=col_idx)
                    if isinstance(value, (int, float)) and "%" not in str(value):
                        value_cell.value = value
                    else:
                        value_cell.value = str(value)
                    value_cell.font = Font(name="Arial", size=16, bold=True)
                    if col_idx == 2:  # Пройдено
                        value_cell.font = Font(name="Arial", size=16, bold=True, color="28a745")
                    elif col_idx == 3:  # Не пройдено
                        value_cell.font = Font(name="Arial", size=16, bold=True, color="dc3545")
                    elif col_idx == 4:  # Успешность
                        value_cell.font = Font(name="Arial", size=16, bold=True, color="4CAF50")
                    value_cell.fill = summary_fill
                    value_cell.alignment = Alignment(horizontal="center", vertical="center")
                    value_cell.border = thin_border

                row += 2

                # Примечание
                non_boolean_count = total_criteria - boolean_criteria_count
                if non_boolean_count > 0:
                    note_cell = ws.merge_cells(f"A{row}:D{row}")
                    note_cell = ws[f"A{row}"]
                    note_cell.value = (
                        f"Примечание: {non_boolean_count} критериев (текстовые и числовые) "
                        "не учитываются в статистике, так как их невозможно оценить "
                        "по принципу верно/неверно."
                    )
                    note_cell.font = Font(name="Arial", size=10, italic=True, color="7f8c8d")
                    note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    row += 1
            else:
                # Только общее количество
                cell = ws.cell(row=row, column=1)
                cell.value = "Всего критериев"
                cell.font = Font(name="Arial", size=9)
                cell.fill = summary_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

                ws.merge_cells(f"B{row}:D{row}")
                value_cell = ws[f"B{row}"]
                value_cell.value = total_criteria
                value_cell.font = Font(name="Arial", size=16, bold=True)
                value_cell.fill = summary_fill
                value_cell.alignment = Alignment(horizontal="center", vertical="center")
                value_cell.border = thin_border

                row += 1

                # Примечание
                ws.merge_cells(f"A{row}:D{row}")
                note_cell = ws[f"A{row}"]
                note_cell.value = (
                    "Примечание: в данной оценке используются только текстовые и числовые критерии, "
                    "которые невозможно оценить по принципу верно/неверно."
                )
                note_cell.font = Font(name="Arial", size=10, italic=True, color="7f8c8d")
                note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                row += 1

            row += 1

            # Футер
            ws.merge_cells(f"A{row}:D{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = "Автоматически сгенерированный отчет"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

            ws.merge_cells(f"A{row}:D{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = f"ID оценки: #{evaluation_id}"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")

            # Настройка ширины столбцов
            ws.column_dimensions["A"].width = 5
            ws.column_dimensions["B"].width = 40
            ws.column_dimensions["C"].width = 20
            ws.column_dimensions["D"].width = 40

            # Сохраняем файл
            wb.save(str(excel_path))

            logger.info(f"Excel report generated: {excel_path}")
            return excel_path

        except Exception as e:
            logger.error(f"Error generating Excel report: {str(e)}", exc_info=True)
            raise

    def generate_employee_report(
        self,
        employee_id: int,
        employee_name: str,
        organization_name: str,
        position: Optional[str] = None,
        phone: Optional[str] = None,
        telegram_id: Optional[int] = None,
        username: Optional[str] = None,
        hire_date: Optional[str] = None,
        is_active: bool = True,
        total_evaluations: int = 0,
        avg_score: Optional[float] = None,
        best_score: Optional[float] = None,
        worst_score: Optional[float] = None,
        last_evaluation_date: Optional[str] = None,
    ) -> Path:
        """Generate Excel report for employee.
        
        Args:
            employee_id: Employee ID.
            employee_name: Employee full name.
            organization_name: Organization name.
            position: Employee position.
            phone: Employee phone.
            telegram_id: Employee Telegram ID.
            username: Employee username.
            hire_date: Employee hire date.
            is_active: Whether employee is active.
            total_evaluations: Total number of evaluations.
            avg_score: Average score.
            best_score: Best score.
            worst_score: Worst score.
            last_evaluation_date: Last evaluation date.
            
        Returns:
            Path to generated Excel file.
        """
        try:
            temp_dir = Path(tempfile.gettempdir())
            excel_filename = f"employee_report_{employee_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            excel_path = temp_dir / excel_filename

            wb = Workbook()
            ws = wb.active
            ws.title = "Отчет по сотруднику"

            title_font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
            header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            normal_font = Font(name="Arial", size=11)
            bold_font = Font(name="Arial", size=11, bold=True)

            title_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            light_gray_fill = PatternFill(start_color="f9f9f9", end_color="f9f9f9", fill_type="solid")
            summary_fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")

            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin")
            )
            thick_border = Border(
                left=Side(style="thick"), right=Side(style="thick"),
                top=Side(style="thick"), bottom=Side(style="thick")
            )

            row = 1

            # Заголовок
            ws.merge_cells(f"A{row}:B{row}")
            cell = ws[f"A{row}"]
            cell.value = "Отчет по сотруднику"
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thick_border
            ws.row_dimensions[row].height = 30
            row += 1

            ws.merge_cells(f"A{row}:B{row}")
            cell = ws[f"A{row}"]
            cell.value = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            cell.font = Font(name="Arial", size=10, color="7f8c8d")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 2

            # Информация о сотруднике
            info_data = [
                ["Организация:", organization_name],
                ["ФИО:", employee_name],
            ]
            if position:
                info_data.append(["Должность:", position])
            if phone:
                info_data.append(["Телефон:", phone])
            if telegram_id:
                info_data.append(["Telegram ID:", str(telegram_id)])
            if username:
                info_data.append(["Username:", username])
            if hire_date:
                info_data.append(["Дата найма:", hire_date])
            info_data.append(["Статус:", "Активен" if is_active else "Неактивен"])

            for label, value in info_data:
                ws[f"A{row}"] = label
                ws[f"A{row}"].font = bold_font
                ws[f"A{row}"].fill = light_gray_fill
                ws[f"A{row}"].border = thin_border
                ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")

                ws[f"B{row}"] = value
                ws[f"B{row}"].font = normal_font
                ws[f"B{row}"].fill = light_gray_fill
                ws[f"B{row}"].border = thin_border
                ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="center")
                row += 1

            row += 1

            # Статистика
            stats_title = ws.cell(row=row, column=1)
            stats_title.value = "Статистика по замерам"
            stats_title.font = Font(name="Arial", size=14, bold=True, color="2c3e50")
            row += 1

            stats_headers = ["Всего замеров"]
            stats_values = [total_evaluations]

            if avg_score is not None:
                stats_headers.append("Средний балл")
                stats_values.append(f"{avg_score:.1f}%")
            if best_score is not None:
                stats_headers.append("Лучший результат")
                stats_values.append(f"{best_score:.1f}%")
            if worst_score is not None:
                stats_headers.append("Худший результат")
                stats_values.append(f"{worst_score:.1f}%")

            for col_idx, (header, value) in enumerate(zip(stats_headers, stats_values), 1):
                cell = ws.cell(row=row, column=col_idx)
                cell.value = str(header)
                cell.font = Font(name="Arial", size=9)
                cell.fill = summary_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

                value_cell = ws.cell(row=row + 1, column=col_idx)
                if "%" not in str(value):
                    try:
                        value_cell.value = int(value)
                    except ValueError:
                        value_cell.value = str(value)
                else:
                    value_cell.value = str(value)
                
                if col_idx == 2 and avg_score is not None:  # Средний балл
                    value_cell.font = Font(name="Arial", size=16, bold=True, color="4CAF50")
                elif col_idx == 3 and best_score is not None:  # Лучший результат
                    value_cell.font = Font(name="Arial", size=16, bold=True, color="28a745")
                elif col_idx == 4 and worst_score is not None:  # Худший результат
                    value_cell.font = Font(name="Arial", size=16, bold=True, color="dc3545")
                else:
                    value_cell.font = Font(name="Arial", size=16, bold=True)
                
                value_cell.fill = summary_fill
                value_cell.alignment = Alignment(horizontal="center", vertical="center")
                value_cell.border = thin_border

            row += 2

            if last_evaluation_date:
                note_cell = ws.merge_cells(f"A{row}:B{row}")
                note_cell = ws[f"A{row}"]
                note_cell.value = f"Последний замер: {last_evaluation_date}"
                note_cell.font = Font(name="Arial", size=10, italic=True, color="7f8c8d")
                note_cell.alignment = Alignment(horizontal="left", vertical="center")
                row += 1

            row += 1

            # Футер
            ws.merge_cells(f"A{row}:B{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = "Автоматически сгенерированный отчет"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

            ws.merge_cells(f"A{row}:B{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = f"ID сотрудника: #{employee_id}"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 40

            wb.save(str(excel_path))
            logger.info(f"Excel employee report generated: {excel_path}")
            return excel_path

        except Exception as e:
            logger.error(f"Error generating employee Excel report: {str(e)}", exc_info=True)
            raise

    def generate_criterion_report(
        self,
        criterion_id: int,
        criterion_name: str,
        organization_name: str,
        code: str,
        value_type: str,
        evaluation_type_id: int,
        category_id: Optional[int] = None,
        is_required: bool = False,
        is_active: bool = True,
        description: Optional[str] = None,
        sort_order: int = 0,
        usage_count: int = 0,
        criterion_sets: Optional[List[str]] = None,
    ) -> Path:
        """Generate Excel report for criterion."""
        try:
            temp_dir = Path(tempfile.gettempdir())
            excel_filename = f"criterion_report_{criterion_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            excel_path = temp_dir / excel_filename

            wb = Workbook()
            ws = wb.active
            ws.title = "Отчет по критерию"

            title_font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
            normal_font = Font(name="Arial", size=11)
            bold_font = Font(name="Arial", size=11, bold=True)

            title_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            light_gray_fill = PatternFill(start_color="f9f9f9", end_color="f9f9f9", fill_type="solid")
            summary_fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")

            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin")
            )
            thick_border = Border(
                left=Side(style="thick"), right=Side(style="thick"),
                top=Side(style="thick"), bottom=Side(style="thick")
            )

            row = 1

            # Заголовок
            ws.merge_cells(f"A{row}:B{row}")
            cell = ws[f"A{row}"]
            cell.value = "Отчет по критерию"
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thick_border
            ws.row_dimensions[row].height = 30
            row += 1

            ws.merge_cells(f"A{row}:B{row}")
            cell = ws[f"A{row}"]
            cell.value = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            cell.font = Font(name="Arial", size=10, color="7f8c8d")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 2

            # Информация о критерии
            info_data = [
                ["Организация:", organization_name],
                ["Название:", criterion_name],
                ["Код:", code],
                ["Тип значения:", value_type],
                ["Тип оценки ID:", str(evaluation_type_id)],
                ["Категория ID:", str(category_id) if category_id else "Нет"],
                ["Обязательный:", "Да" if is_required else "Нет"],
                ["Активен:", "Да" if is_active else "Нет"],
                ["Порядок сортировки:", str(sort_order)],
            ]
            if description:
                info_data.append(["Описание:", description])

            for label, value in info_data:
                ws[f"A{row}"] = label
                ws[f"A{row}"].font = bold_font
                ws[f"A{row}"].fill = light_gray_fill
                ws[f"A{row}"].border = thin_border
                ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")

                ws[f"B{row}"] = value
                ws[f"B{row}"].font = normal_font
                ws[f"B{row}"].fill = light_gray_fill
                ws[f"B{row}"].border = thin_border
                ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                row += 1

            row += 1

            # Статистика использования
            stats_title = ws.cell(row=row, column=1)
            stats_title.value = "Статистика использования"
            stats_title.font = Font(name="Arial", size=14, bold=True, color="2c3e50")
            row += 1

            cell = ws.cell(row=row, column=1)
            cell.value = "Использован в замерах"
            cell.font = Font(name="Arial", size=9)
            cell.fill = summary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            value_cell = ws.cell(row=row + 1, column=1)
            value_cell.value = usage_count
            value_cell.font = Font(name="Arial", size=16, bold=True)
            value_cell.fill = summary_fill
            value_cell.alignment = Alignment(horizontal="center", vertical="center")
            value_cell.border = thin_border

            row += 2

            if criterion_sets:
                ws.merge_cells(f"A{row}:B{row}")
                note_cell = ws[f"A{row}"]
                note_cell.value = f"Входит в наборы ({len(criterion_sets)}):"
                note_cell.font = Font(name="Arial", size=11, bold=True)
                note_cell.alignment = Alignment(horizontal="left", vertical="center")
                row += 1

                for set_name in criterion_sets[:20]:
                    ws.merge_cells(f"A{row}:B{row}")
                    note_cell = ws[f"A{row}"]
                    note_cell.value = f"  • {set_name}"
                    note_cell.font = Font(name="Arial", size=10)
                    note_cell.alignment = Alignment(horizontal="left", vertical="center")
                    row += 1

                if len(criterion_sets) > 20:
                    ws.merge_cells(f"A{row}:B{row}")
                    note_cell = ws[f"A{row}"]
                    note_cell.value = f"  ... и еще {len(criterion_sets) - 20}"
                    note_cell.font = Font(name="Arial", size=10, italic=True)
                    note_cell.alignment = Alignment(horizontal="left", vertical="center")
                    row += 1
            else:
                ws.merge_cells(f"A{row}:B{row}")
                note_cell = ws[f"A{row}"]
                note_cell.value = "Входит в наборы: Нет"
                note_cell.font = Font(name="Arial", size=10)
                note_cell.alignment = Alignment(horizontal="left", vertical="center")
                row += 1

            row += 1

            # Футер
            ws.merge_cells(f"A{row}:B{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = "Автоматически сгенерированный отчет"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

            ws.merge_cells(f"A{row}:B{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = f"ID критерия: #{criterion_id}"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 40

            wb.save(str(excel_path))
            logger.info(f"Excel criterion report generated: {excel_path}")
            return excel_path

        except Exception as e:
            logger.error(f"Error generating criterion Excel report: {str(e)}", exc_info=True)
            raise

    def generate_criterion_set_report(
        self,
        criterion_set_id: int,
        criterion_set_name: str,
        organization_name: str,
        is_default: bool = False,
        is_active: bool = True,
        description: Optional[str] = None,
        criteria: Optional[List[Dict]] = None,
        usage_count: int = 0,
    ) -> Path:
        """Generate Excel report for criterion set."""
        try:
            temp_dir = Path(tempfile.gettempdir())
            excel_filename = f"criterion_set_report_{criterion_set_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            excel_path = temp_dir / excel_filename

            wb = Workbook()
            ws = wb.active
            ws.title = "Отчет по набору критериев"

            title_font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
            header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
            normal_font = Font(name="Arial", size=11)
            bold_font = Font(name="Arial", size=11, bold=True)

            title_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            header_fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
            light_gray_fill = PatternFill(start_color="f9f9f9", end_color="f9f9f9", fill_type="solid")
            summary_fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")

            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin")
            )
            thick_border = Border(
                left=Side(style="thick"), right=Side(style="thick"),
                top=Side(style="thick"), bottom=Side(style="thick")
            )

            row = 1

            # Заголовок
            ws.merge_cells(f"A{row}:D{row}")
            cell = ws[f"A{row}"]
            cell.value = "Отчет по набору критериев"
            cell.font = title_font
            cell.fill = title_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thick_border
            ws.row_dimensions[row].height = 30
            row += 1

            ws.merge_cells(f"A{row}:D{row}")
            cell = ws[f"A{row}"]
            cell.value = f"Дата формирования: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
            cell.font = Font(name="Arial", size=10, color="7f8c8d")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 2

            # Информация о наборе
            info_data = [
                ["Организация:", organization_name],
                ["Название:", criterion_set_name],
                ["Дефолтный:", "Да" if is_default else "Нет"],
                ["Активен:", "Да" if is_active else "Нет"],
            ]
            if description:
                info_data.append(["Описание:", description])

            for label, value in info_data:
                ws[f"A{row}"] = label
                ws[f"A{row}"].font = bold_font
                ws[f"A{row}"].fill = light_gray_fill
                ws[f"A{row}"].border = thin_border
                ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")

                ws[f"B{row}"] = value
                ws[f"B{row}"].font = normal_font
                ws[f"B{row}"].fill = light_gray_fill
                ws[f"B{row}"].border = thin_border
                ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                ws.merge_cells(f"B{row}:D{row}")
                row += 1

            row += 1

            # Критерии в наборе
            if criteria:
                criteria_title = ws.cell(row=row, column=1)
                criteria_title.value = f"Критерии в наборе ({len(criteria)})"
                criteria_title.font = Font(name="Arial", size=14, bold=True, color="2c3e50")
                row += 1

                headers = ["№", "Название", "Код", "Тип"]
                for col_idx, header in enumerate(headers, 1):
                    cell = ws.cell(row=row, column=col_idx)
                    cell.value = header
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.border = thick_border
                ws.row_dimensions[row].height = 25
                row += 1

                for idx, criterion in enumerate(criteria[:100], 1):
                    ws.cell(row=row, column=1).value = idx
                    ws.cell(row=row, column=1).font = normal_font
                    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center", vertical="center")
                    ws.cell(row=row, column=1).border = thin_border

                    ws.cell(row=row, column=2).value = criterion.get('name', 'N/A')
                    ws.cell(row=row, column=2).font = normal_font
                    ws.cell(row=row, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws.cell(row=row, column=2).border = thin_border

                    ws.cell(row=row, column=3).value = criterion.get('code', 'N/A')
                    ws.cell(row=row, column=3).font = normal_font
                    ws.cell(row=row, column=3).alignment = Alignment(horizontal="left", vertical="center")
                    ws.cell(row=row, column=3).border = thin_border

                    ws.cell(row=row, column=4).value = criterion.get('value_type', 'N/A')
                    ws.cell(row=row, column=4).font = normal_font
                    ws.cell(row=row, column=4).alignment = Alignment(horizontal="left", vertical="center")
                    ws.cell(row=row, column=4).border = thin_border

                    if idx % 2 == 0:
                        for col in range(1, 5):
                            ws.cell(row=row, column=col).fill = light_gray_fill

                    row += 1

                if len(criteria) > 100:
                    ws.merge_cells(f"A{row}:D{row}")
                    note_cell = ws[f"A{row}"]
                    note_cell.value = f"... и еще {len(criteria) - 100} критериев"
                    note_cell.font = Font(name="Arial", size=10, italic=True, color="7f8c8d")
                    note_cell.alignment = Alignment(horizontal="left", vertical="center")
                    row += 1

                row += 1
            else:
                ws.merge_cells(f"A{row}:D{row}")
                note_cell = ws[f"A{row}"]
                note_cell.value = "Критерии в наборе: Нет"
                note_cell.font = Font(name="Arial", size=11)
                note_cell.alignment = Alignment(horizontal="left", vertical="center")
                row += 2

            # Статистика использования
            stats_title = ws.cell(row=row, column=1)
            stats_title.value = "Статистика использования"
            stats_title.font = Font(name="Arial", size=14, bold=True, color="2c3e50")
            row += 1

            cell = ws.cell(row=row, column=1)
            cell.value = "Использован в замерах"
            cell.font = Font(name="Arial", size=9)
            cell.fill = summary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

            value_cell = ws.cell(row=row + 1, column=1)
            value_cell.value = usage_count
            value_cell.font = Font(name="Arial", size=16, bold=True)
            value_cell.fill = summary_fill
            value_cell.alignment = Alignment(horizontal="center", vertical="center")
            value_cell.border = thin_border

            row += 3

            # Футер
            ws.merge_cells(f"A{row}:D{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = "Автоматически сгенерированный отчет"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")
            row += 1

            ws.merge_cells(f"A{row}:D{row}")
            footer_cell = ws[f"A{row}"]
            footer_cell.value = f"ID набора: #{criterion_set_id}"
            footer_cell.font = Font(name="Arial", size=9, color="7f8c8d")
            footer_cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.column_dimensions["A"].width = 5
            ws.column_dimensions["B"].width = 40
            ws.column_dimensions["C"].width = 15
            ws.column_dimensions["D"].width = 15

            wb.save(str(excel_path))
            logger.info(f"Excel criterion set report generated: {excel_path}")
            return excel_path

        except Exception as e:
            logger.error(f"Error generating criterion set Excel report: {str(e)}", exc_info=True)
            raise

