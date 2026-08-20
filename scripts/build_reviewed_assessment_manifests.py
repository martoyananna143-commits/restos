"""Build reviewed Wave A manifests from exact read-only Excel sources.

This development command never writes to the source workbooks. It deliberately
extracts only methodology rows (section, prompt, and weight), never historical
names, dates, answers, or comments.
"""

from __future__ import annotations

import argparse
from hashlib import sha256
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


OUTPUT = (
    Path(__file__).parents[1] / "app/internal/data/assessment_template_import/manifests"
)
SOURCES = {
    "cook-kln": (
        "КЛН повара_Камелот.xlsx",
        "6bdcdcb35a6003769552209e6439bb6c56f9ac968120da6bdeb650db592e064e",
        4,
        "КЛН повара",
    ),
    "waiter-kln": (
        ("НОВЫЙ ИСПРАВЛЕННЫЙ " "КЛН-корректировка официант - Камелот, новый .xlsx"),
        "e0e34ffcb0304fc3db9f0c9b790acc8613d54cc76488ef1b2bba6d08b584048a",
        5,
        "КЛН официанта",
    ),
    "hostess-kln": (
        ("НОВЫЙ ИСПРАВЛЕННЫЙ " "КЛН-корректировка хостес .xlsx"),
        "065b9e74097bb7cdda55d3d36df7840b89d5c91bcaef226ab5f6abfabef80938",
        5,
        "КЛН хостес",
    ),
    "itci-kln": (
        "SM, КЛН_ITCI,  Камелот - новый .xlsx",
        "43d2993f09a616ead491b412ef48ccc778392f786d9d4f45da17d4b59aaa33d9",
        5,
        "КЛН ITCI",
    ),
    "kitchen-practicum": (
        "КЛН Руководителей_Камелот.xlsx",
        "b19824ecac962517df729aa7270aca97ce77a191f211136e3e10007660e63c16",
        5,
        "Практикум кухни",
    ),
    "bar-practicum": (
        "КЛН Руководителей_Камелот.xlsx",
        "b19824ecac962517df729aa7270aca97ce77a191f211136e3e10007660e63c16",
        7,
        "Практикум бара",
    ),
}
SERVICE_SOURCE = (
    "SM, Бланк оценки сервиса ресторана, Камелот - новый .xlsx",
    "13b24f5ebf320d7bc7ed2c0313dd99b8b0f44c8408ce688c7e3406f47dea8e4d",
)
PRODUCTION_SOURCE = (
    "БО производства WP_Дебри.xlsx",
    "933fca2cb9b9b8b043e2085c547193bc9af6b2fefbe2392f2abe1537238c15bd",
)
PRODUCTION_METRICS = {
    "Скорость": "speed",
    "Порядок": "order",
    "Вкус": "taste",
    "Люди": "people",
    "Пространство": "space",
    "Экономика": "economics",
}
WAITER_EQUAL_WEIGHT = "1.0"
WAITER_WEIGHT_POLICY = "owner_equal_criterion_weights_v1"


def digest(path: Path) -> str:
    value = sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def slug(value: str, fallback: str) -> str:
    transliteration = str.maketrans(
        {
            "а": "a",
            "б": "b",
            "в": "v",
            "г": "g",
            "д": "d",
            "е": "e",
            "ё": "e",
            "ж": "zh",
            "з": "z",
            "и": "i",
            "й": "y",
            "к": "k",
            "л": "l",
            "м": "m",
            "н": "n",
            "о": "o",
            "п": "p",
            "р": "r",
            "с": "s",
            "т": "t",
            "у": "u",
            "ф": "f",
            "х": "h",
            "ц": "c",
            "ч": "ch",
            "ш": "sh",
            "щ": "sch",
            "ъ": "",
            "ы": "y",
            "ь": "",
            "э": "e",
            "ю": "yu",
            "я": "ya",
        }
    )
    candidate = value.lower().translate(transliteration)
    candidate = re.sub(r"[^a-z0-9]+", "-", candidate).strip("-")
    return candidate[:90].rstrip("-") or fallback


def mappings(
    section: str, prompt: str, weight: str | None
) -> list[dict[str, str | None]]:
    text = f"{section} {prompt}".lower()
    codes = ["people"]
    if any(
        value in text for value in ("гост", "сервис", "бронир", "официант", "хостес")
    ):
        codes.append("service")
    if any(value in text for value in ("вкус", "сырь", "блюд", "напит", "технолог")):
        codes.append("taste")
    if any(value in text for value in ("скорост", "минут", "вовремя", "время", "гудк")):
        codes.append("speed")
    if any(
        value in text
        for value in ("поряд", "чист", "инвентар", "оборуд", "чек-лист", "ротац")
    ):
        codes.append("order")
    if any(
        value in text
        for value in (
            "санитар",
            "гигиен",
            "сроки годности",
            "товарное соседство",
            "бракераж",
        )
    ):
        codes.append("food_safety")
    if any(value in text for value in ("выруч", "продаж", "эконом", "go-list")):
        codes.append("economics")
    if any(
        value in text for value in ("место проведения", "рабочем месте", "зала", "стол")
    ):
        codes.append("space")
    return [
        {"code": code, "weight": weight, "direction": "positive"}
        for code in dict.fromkeys(codes)
    ]


def extract(
    code: str, path: Path, expected_sha: str, sheet_index: int, name: str
) -> dict[str, Any]:
    if digest(path) != expected_sha:
        raise RuntimeError(f"source sha mismatch: {path.name}")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.worksheets[sheet_index - 1]
        header_row = None
        prompt_column = None
        weight_column = None
        section_column = None
        for row_index, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=80, max_col=8, values_only=True),
            start=1,
        ):
            for column, value in enumerate(row, start=1):
                if isinstance(value, str) and "Оцениваемые показатели" in value:
                    header_row = row_index
                    prompt_column = column
                    section_column = column - 1 if column > 1 else None
                    weight_column = column + 1
                    break
            if header_row is not None:
                break
        if header_row is None or prompt_column is None or weight_column is None:
            raise RuntimeError(f"methodology header not found: {path.name}")

        sections: list[dict[str, Any]] = []
        by_code: dict[str, dict[str, Any]] = {}
        current_section = name if section_column is None else "Общие критерии"
        used_item_codes: set[str] = set()
        for row in sheet.iter_rows(
            min_row=header_row + 1,
            max_row=100,
            max_col=max(weight_column, 3),
            values_only=True,
        ):
            prompt = row[prompt_column - 1]
            raw_weight = row[weight_column - 1]
            section_value = (
                row[section_column - 1] if section_column is not None else None
            )
            if not isinstance(prompt, str) or not prompt.strip():
                continue
            prompt = " ".join(prompt.split())
            if (
                prompt.startswith("MAX СУММА")
                or prompt.startswith("Результат")
                or prompt.startswith("Средний показатель")
                or prompt.startswith("Количество замеров")
                or prompt.startswith("©")
                or isinstance(raw_weight, str)
            ):
                continue
            if isinstance(section_value, str) and section_value.strip():
                current_section = " ".join(section_value.split())
            section_code = slug(current_section, "section")
            if section_code not in by_code:
                section = {
                    "code": section_code,
                    "title": current_section,
                    "description": None,
                    "kind": "section",
                    "weight": None,
                    "items": [],
                }
                by_code[section_code] = section
                sections.append(section)
            item_code = slug(prompt, f"criterion-{len(used_item_codes) + 1}")
            base = item_code
            suffix = 2
            while item_code in used_item_codes:
                item_code = f"{base[:85].rstrip('-')}-{suffix}"
                suffix += 1
            used_item_codes.add(item_code)
            source_weight = (
                str(raw_weight) if isinstance(raw_weight, (int, float)) else None
            )
            weight = WAITER_EQUAL_WEIGHT if code == "waiter-kln" else source_weight
            by_code[section_code]["items"].append(
                {
                    "code": item_code,
                    "prompt": prompt,
                    "guidance": None,
                    "response_type": "boolean",
                    "required": True,
                    "weight": weight,
                    "min_value": None,
                    "max_value": None,
                    "passing_value": None,
                    "evidence_mode": "optional_comment",
                    "criticality": "normal",
                    "config": {},
                    "options": [],
                    "metrics": mappings(current_section, prompt, weight),
                }
            )
    finally:
        workbook.close()
    return {
        "schema_version": 1,
        "source": {
            "filename": path.name,
            "sha256": expected_sha,
            "provenance": "read_only_methodology_source",
        },
        "template": {
            "code": code,
            "name": name,
            "activity_type": "evaluation",
            "description": (
                "КЛН официанта: 43 критерия с одинаковым весом 1."
                if code == "waiter-kln"
                else "Company-private draft imported from reviewed methodology source."
            ),
            "methodology": {
                "code": f"{code}-methodology",
                "title": name,
                "body": (
                    "Оценка по 43 критериям Да/Нет. Все критерии имеют "
                    "одинаковый вес 1 согласно утверждённому продуктовому решению. "
                    "Отсутствующие исторические результаты и персональные данные "
                    "не импортируются."
                    if code == "waiter-kln"
                    else "Взвешенная оценка по критериям Да/Нет. Отсутствующие "
                    "исторические результаты и персональные данные не импортируются."
                ),
                "version": 1,
            },
            "scoring": {
                "algorithm": "weighted_v1",
                "version": 1,
                "config": {
                    "rounding": "half_up",
                    "score_scale": 100,
                    **(
                        {
                            "weight_policy": WAITER_WEIGHT_POLICY,
                            "criterion_weight": WAITER_EQUAL_WEIGHT,
                            "source_weights_used": False,
                        }
                        if code == "waiter-kln"
                        else {}
                    ),
                },
            },
            "sections": sections,
        },
    }


def item_document(
    *,
    code: str,
    prompt: str,
    guidance: str | None,
    weight: str,
    metric_codes: list[str],
    criticality: str = "normal",
) -> dict[str, Any]:
    return {
        "code": code,
        "prompt": prompt,
        "guidance": guidance,
        "response_type": "boolean",
        "required": True,
        "weight": weight,
        "min_value": None,
        "max_value": None,
        "passing_value": None,
        "evidence_mode": "optional_comment",
        "criticality": criticality,
        "config": {"comment_on_negative": True},
        "options": [],
        "metrics": [
            {"code": metric, "weight": weight, "direction": "positive"}
            for metric in dict.fromkeys(metric_codes)
        ],
    }


def walkthrough_document(
    *,
    code: str,
    name: str,
    source_path: Path,
    expected_sha: str,
    sections: list[dict[str, Any]],
) -> dict[str, Any]:
    if digest(source_path) != expected_sha:
        raise RuntimeError(f"source sha mismatch: {source_path.name}")
    return {
        "schema_version": 1,
        "source": {
            "filename": source_path.name,
            "sha256": expected_sha,
            "provenance": "read_only_methodology_source",
        },
        "template": {
            "code": code,
            "name": name,
            "activity_type": "walkthrough",
            "description": "Company-private operational walkthrough from reviewed methodology source.",
            "methodology": {
                "code": f"{code}-methodology",
                "title": name,
                "body": "Каждое завершённое заполнение является отдельным обходом. Исторические ответы, даты, комментарии и персональные данные не импортируются.",
                "version": 1,
            },
            "scoring": {
                "algorithm": "weighted_v1",
                "version": 1,
                "config": {"rounding": "half_up", "score_scale": 100},
            },
            "sections": sections,
        },
    }


def extract_service_walkthrough(path: Path, expected_sha: str) -> dict[str, Any]:
    if digest(path) != expected_sha:
        raise RuntimeError(f"source sha mismatch: {path.name}")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook.worksheets[0]
        phase = "Service Design"
        group = "Общие критерии"
        sections: list[dict[str, Any]] = []
        by_code: dict[str, dict[str, Any]] = {}
        used: set[str] = set()
        for row in sheet.iter_rows(values_only=True):
            title = row[0] if row else None
            block = row[1] if len(row) > 1 else None
            weight = row[2] if len(row) > 2 else None
            if not isinstance(title, str) or not title.strip():
                continue
            title = " ".join(title.split())
            upper = title.upper()
            if upper in {"SERVICE DESIGN", "SERVICE MANAGEMENT"}:
                phase = title.title()
                group = "Общие критерии"
                continue
            if not isinstance(weight, (int, float)):
                if isinstance(block, str) and block.startswith("=SUM("):
                    group = title
                continue
            section_title = f"{phase} · {group}"
            section_code = slug(section_title, "service-section")
            if section_code not in by_code:
                section = {
                    "code": section_code,
                    "title": section_title,
                    "description": None,
                    "kind": "section",
                    "weight": None,
                    "items": [],
                }
                by_code[section_code] = section
                sections.append(section)
            item_code = unique_slug(title, used)
            text = f"{block or ''} {title}".lower()
            metric_codes = ["service"]
            if any(
                value in text for value in ("скорост", "время", "минут", "своевременно")
            ):
                metric_codes.append("speed")
            if any(
                value in text
                for value in (
                    "чист",
                    "поряд",
                    "чек-лист",
                    "срок годности",
                    "fifo",
                    "маркиров",
                )
            ):
                metric_codes.append("order")
            if any(
                value in text
                for value in (
                    "сотруд",
                    "знает",
                    "обуч",
                    "менеджер",
                    "официант",
                    "хостес",
                )
            ):
                metric_codes.append("people")
            if any(
                value in text
                for value in ("простран", "зал", "уборн", "оборуд", "раздевал")
            ):
                metric_codes.append("space")
            if any(
                value in text
                for value in ("санитар", "срок годности", "товарного соседства", "fifo")
            ):
                metric_codes.append("food_safety")
            by_code[section_code]["items"].append(
                item_document(
                    code=item_code,
                    prompt=title,
                    guidance=None,
                    weight=str(weight),
                    metric_codes=metric_codes,
                )
            )
    finally:
        workbook.close()
    if sum(len(section["items"]) for section in sections) != 117:
        raise RuntimeError("service walkthrough criterion count mismatch")
    return walkthrough_document(
        code="restaurant-service-walkthrough",
        name="Оценка сервиса ресторана",
        source_path=path,
        expected_sha=expected_sha,
        sections=sections,
    )


def extract_production_walkthrough(path: Path, expected_sha: str) -> dict[str, Any]:
    if digest(path) != expected_sha:
        raise RuntimeError(f"source sha mismatch: {path.name}")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook.worksheets[1]
        zone = "Общие зоны"
        sections: list[dict[str, Any]] = []
        by_code: dict[str, dict[str, Any]] = {}
        used: set[str] = set()
        for row in sheet.iter_rows(values_only=True):
            prompt = row[0] if row else None
            guidance = row[1] if len(row) > 1 else None
            category = row[3] if len(row) > 3 else None
            weight = row[4] if len(row) > 4 else None
            if (
                isinstance(prompt, str)
                and prompt.strip()
                and category not in PRODUCTION_METRICS
            ):
                if all(value in (None, "") for value in row[1:5]):
                    zone = " ".join(prompt.split())
                continue
            if (
                not isinstance(prompt, str)
                or category not in PRODUCTION_METRICS
                or not isinstance(weight, (int, float))
            ):
                continue
            prompt = " ".join(prompt.split())
            section_code = slug(zone, "production-zone")
            if section_code not in by_code:
                section = {
                    "code": section_code,
                    "title": zone,
                    "description": None,
                    "kind": "zone",
                    "weight": None,
                    "items": [],
                }
                by_code[section_code] = section
                sections.append(section)
            metric_codes = [PRODUCTION_METRICS[category]]
            lower = f"{prompt} {guidance or ''}".lower()
            if any(
                value in lower
                for value in (
                    "санпин",
                    "пищев",
                    "дезинф",
                    "хассп",
                    "haccp",
                    "срок годности",
                    "товарн",
                )
            ):
                metric_codes.append("food_safety")
            critical = "критическое нарушение" in lower
            by_code[section_code]["items"].append(
                item_document(
                    code=unique_slug(prompt, used),
                    prompt=prompt,
                    guidance=(
                        " ".join(guidance.split())
                        if isinstance(guidance, str) and guidance.strip()
                        else None
                    ),
                    weight=str(weight),
                    metric_codes=metric_codes,
                    criticality="stop_factor" if critical else "normal",
                )
            )
    finally:
        workbook.close()
    if sum(len(section["items"]) for section in sections) != 199:
        raise RuntimeError("production walkthrough criterion count mismatch")
    return walkthrough_document(
        code="production-walkthrough",
        name="Производственный обход",
        source_path=path,
        expected_sha=expected_sha,
        sections=sections,
    )


def unique_slug(value: str, used: set[str]) -> str:
    candidate = slug(value, f"criterion-{len(used) + 1}")
    base = candidate
    suffix = 2
    while candidate in used:
        candidate = f"{base[:85].rstrip('-')}-{suffix}"
        suffix += 1
    used.add(candidate)
    return candidate


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True, type=Path)
    values = parser.parse_args()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    counts = {}
    for code, (filename, expected_sha, sheet_index, name) in SOURCES.items():
        path = values.source_root / filename
        document = extract(code, path, expected_sha, sheet_index, name)
        target = OUTPUT / f"{code}.json"
        target.write_text(
            json.dumps(document, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        counts[code] = sum(
            len(section["items"]) for section in document["template"]["sections"]
        )
    for document in (
        extract_service_walkthrough(
            values.source_root / SERVICE_SOURCE[0], SERVICE_SOURCE[1]
        ),
        extract_production_walkthrough(
            values.source_root / PRODUCTION_SOURCE[0], PRODUCTION_SOURCE[1]
        ),
    ):
        code = document["template"]["code"]
        (OUTPUT / f"{code}.json").write_text(
            json.dumps(document, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        counts[code] = sum(
            len(section["items"]) for section in document["template"]["sections"]
        )
    print(json.dumps(counts, sort_keys=True))


if __name__ == "__main__":
    main()
